import asyncio
from io import BytesIO

import pytest
from docx import Document
from fastapi import HTTPException
from PIL import Image, ImageDraw
import pytesseract
from starlette.datastructures import Headers
from starlette.datastructures import UploadFile

from app.certificates import TesseractUnavailable
from app.main import app, certificate_extractor, extract_certificate, extract_resume, health


def upload(name: str, content: bytes, content_type: str) -> UploadFile:
    return UploadFile(BytesIO(content), filename=name, headers=Headers({"content-type": content_type}))


def docx_fixture() -> bytes:
    document = Document()
    document.add_paragraph("SKILLS")
    document.add_paragraph("Python, JS, PostgreSQL")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def pdf_fixture() -> bytes:
    objects = [b"1 0 obj<< /Type /Catalog /Pages 2 0 R>>endobj", b"2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1>>endobj", b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources<< /Font<< /F1 4 0 R>> >> /Contents 5 0 R>>endobj", b"4 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica>>endobj", b"5 0 obj<< /Length 50>>stream\nBT /F1 12 Tf 72 720 Td (Python JavaScript) Tj ET\nendstream endobj"]
    body = b"%PDF-1.4\n"; offsets = [0]
    for obj in objects: offsets.append(len(body)); body += obj + b"\n"
    xref = len(body); body += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    return body + f"trailer<< /Root 1 0 R /Size 6>>\nstartxref\n{xref}\n%%EOF".encode()


def call(file: UploadFile):
    return asyncio.run(extract_resume(file))


def call_certificate(file: UploadFile):
    return asyncio.run(extract_certificate(file))


def test_health():
    result = health()
    assert result["status"] == "UP"
    assert result["modelLoaded"] is False


def test_pdf_fixture_extracts_text_and_aliases():
    result = call(upload("resume.pdf", pdf_fixture(), "application/pdf"))
    skills = {item["skill"] for item in result["matchedSkills"]}
    assert {"python", "javascript"}.issubset(skills)


def test_docx_fixture_extracts_text_and_aliases():
    result = call(upload("resume.docx", docx_fixture(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    assert {item["skill"] for item in result["matchedSkills"]} == {"javascript", "postgresql", "python"}


def test_unsupported_type_is_rejected():
    with pytest.raises(HTTPException) as error:
        call(upload("resume.txt", b"Python", "text/plain"))
    assert error.value.status_code == 415


def test_empty_file_is_rejected():
    with pytest.raises(HTTPException) as error:
        call(upload("resume.pdf", b"", "application/pdf"))
    assert error.value.status_code == 400


def test_oversized_file_is_rejected(monkeypatch):
    monkeypatch.setattr("app.main.MAX_FILE_SIZE", 4)
    with pytest.raises(HTTPException) as error:
        call(upload("resume.pdf", b"12345", "application/pdf"))
    assert error.value.status_code == 413


def certificate_pdf_fixture() -> bytes:
    text = b"BT /F1 12 Tf 72 720 Td (Certificate of Python Development) Tj 0 -24 Td (Issued by Example Academy) Tj ET"
    objects = [b"1 0 obj<< /Type /Catalog /Pages 2 0 R>>endobj", b"2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1>>endobj", b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources<< /Font<< /F1 4 0 R>> >> /Contents 5 0 R>>endobj", b"4 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica>>endobj", b"5 0 obj<< /Length " + str(len(text)).encode() + b">>stream\n" + text + b"\nendstream endobj"]
    body = b"%PDF-1.4\n"; offsets = [0]
    for obj in objects: offsets.append(len(body)); body += obj + b"\n"
    xref = len(body); body += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    return body + f"trailer<< /Root 1 0 R /Size 6>>\nstartxref\n{xref}\n%%EOF".encode()


def test_text_layer_certificate_extracts_metadata_and_skills():
    result = call_certificate(upload("certificate.pdf", certificate_pdf_fixture(), "application/pdf"))
    assert result["certificateTitle"] == "Certificate of Python Development"
    assert result["issuer"] == "Example Academy"
    assert result["detectedSkills"][0]["skill"] == "python"
    assert result["needsReview"] is False


def test_corrupted_certificate_is_rejected():
    with pytest.raises(HTTPException) as error:
        call_certificate(upload("certificate.pdf", b"not a pdf", "application/pdf"))
    assert error.value.status_code == 422


def test_certificate_unsupported_file_is_rejected():
    with pytest.raises(HTTPException) as error:
        call_certificate(upload("certificate.txt", b"Certificate", "text/plain"))
    assert error.value.status_code == 415


def test_ocr_unavailable_returns_503(monkeypatch):
    def unavailable(*_args, **_kwargs):
        raise TesseractUnavailable("Tesseract is unavailable")
    monkeypatch.setattr(certificate_extractor, "extract", unavailable)
    with pytest.raises(HTTPException) as error:
        call_certificate(upload("certificate.png", b"image-bytes", "image/png"))
    assert error.value.status_code == 503


def test_image_certificate_ocr():
    try:
        pytesseract.get_tesseract_version()
    except (OSError, pytesseract.TesseractNotFoundError):
        pytest.skip("System Tesseract is unavailable")
    image = Image.new("RGB", (900, 300), "white")
    ImageDraw.Draw(image).text((30, 40), "Certificate of Python\nIssued by Example Academy", fill="black")
    output = BytesIO(); image.save(output, format="PNG")
    result = call_certificate(upload("certificate.png", output.getvalue(), "image/png"))
    assert result["rawText"]
    assert result["needsReview"] is True

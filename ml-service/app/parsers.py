from abc import ABC, abstractmethod
from io import BytesIO

from docx import Document
from pypdf import PdfReader


class ResumeParser(ABC):
    @abstractmethod
    def parse(self, content: bytes) -> str:
        raise NotImplementedError


class PdfResumeParser(ResumeParser):
    def parse(self, content: bytes) -> str:
        reader = PdfReader(BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()


class DocxResumeParser(ResumeParser):
    def parse(self, content: bytes) -> str:
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_text = [cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip()]
        return "\n".join(paragraphs + table_text).strip()


PARSERS = {".pdf": PdfResumeParser(), ".docx": DocxResumeParser()}
